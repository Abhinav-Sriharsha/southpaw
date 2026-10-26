"""South Paw · tools/career/make_sheet.py
Writes tools/career/roster.xlsx from roster.json (made by make_roster.py): every career fighter with name, look and technique,
plus the fighting styles and the three tournaments.  Run:  python3 tools/career/make_sheet.py"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
HERE=os.path.dirname(os.path.abspath(__file__))
d=json.load(open(os.path.join(HERE,'roster.json')))
F='Arial';HEAD=PatternFill('solid',fgColor='1B1D24');HF=Font(name=F,bold=True,color='FFFFFF',size=10);BF=Font(name=F,size=10);IN=Font(name=F,size=10,color='0000FF')
thin=Side(style='thin',color='D0D0D0');BOX=Border(bottom=thin)
HAIR2={'':'Bald','Hair_Buzzed':'Buzzed','Hair_SimpleParted':'Simple parted','Hair_Long':'Long','Hair_Buns':'Buns','g_crew':'Crew cut','g_flattop':'Flat top','g_afro':'Afro',
 'g_hightop':'High-top fade','g_mohawk':'Mohawk','g_fauxhawk':'Faux hawk','g_manbun':'Man bun','g_cornrows':'Cornrows','g_twists':'Twists','g_slick':'Pompadour'}
BOTTOM={'':'Classic trunks','u_shorts_long':'Long trunks','u_mma':'Fight shorts','u_briefs':'Briefs','u_track':'Track pants'}
FEET={'low':'Low boots','mid':'Classic boots','u_boots':'High boots','barefoot':'Barefoot'}
TOUR={t['id']:t for t in d['tours']};ST=d['styles']
def build(w,m):
  a='Very lean' if w<-0.6 else 'Lean' if w<-0.2 else 'Middleweight' if w<0.2 else 'Heavy' if w<0.6 else 'Heavyweight'
  b='lean' if m<0.2 else 'toned' if m<0.55 else 'built' if m<0.85 else 'bodybuilder';return a+', '+b
def textc(h):h=h.lstrip('#');r,g,b=int(h[:2],16),int(h[2:4],16),int(h[4:],16);return '000000' if (r*299+g*587+b*114)/1000>140 else 'FFFFFF'
def head(ws,cols,widths):
  for i,(c,w) in enumerate(zip(cols,widths),1):
    x=ws.cell(row=1,column=i,value=c);x.font=HF;x.fill=HEAD;x.alignment=Alignment(vertical='center',wrap_text=True);ws.column_dimensions[get_column_letter(i)].width=w
  ws.row_dimensions[1].height=30;ws.freeze_panes='D2'
wb=Workbook();ws=wb.active;ws.title='Roster'
cols=['Tournament','Seed','Name','Nickname','Hometown','Style','Technique','Signature punch','Skill (0-1)','Danger (1-10)','Toughness','Power','Takes damage','Hair','Hair colour','Facial hair','Face','Eyebrows','Skin','Eyes','Build','Shorts','Shorts colour','Trim','Top','Headgear','Gloves','Boots','Sole']
head(ws,cols,[13,6,20,16,12,16,46,16,9,9,9,8,10,16,12,12,15,10,11,11,22,14,12,11,12,10,11,16,8])
order={'amateur':0,'pro':1,'ultimate':2};rows=sorted(d['fighters'],key=lambda f:(order[f['tour']],f['rank']))
for r,f in enumerate(rows,2):
  l=f['look'];n=f['names'];st=ST[f['style']];T=TOUR[f['tour']]
  g=f['look2']
  vals=[T['short'],f['rank'],f['name'],f['nick'],f['from'],st['name'],st['desc'],f['sig'].capitalize(),f['s'],
    f'=MAX(1,MIN(10,ROUND(I{r}*8+IF(A{r}="Amateur",0.5,1),0)))',f['tough'],f['power'],f'=1/K{r}',
    HAIR2[g['hair']],n['hair'] if g['hair'] else '',('Beard' if g['beard'] else 'Clean shaven'),g['face'] or 'Original','Thin' if g['brow']=='Eyebrows_Female' else 'Regular',
    n['skin'],n['eyes'],build(g['fat'],(g['mus']-0.4)/0.55),BOTTOM[g['bottom']],n['trunks'],n['band'],'Vest · '+n['trunks'] if g['top'] else 'Bare chest',
    'Yes' if g['gear'] else 'No',n['wraps'],FEET[g['feet']]+' · '+n['shoe'],n['sole']]
  for c,v in enumerate(vals,1):
    x=ws.cell(row=r,column=c,value=v);x.font=IN if c in (9,11,12) else BF;x.border=BOX;x.alignment=Alignment(vertical='top',wrap_text=c==7)
  ws.cell(row=r,column=9).number_format='0.00';ws.cell(row=r,column=11).number_format='0.00"x"';ws.cell(row=r,column=12).number_format='0.00"x"';ws.cell(row=r,column=13).number_format='0%'
  # colour cells show the colour itself
  for c,hx in [(15,g['hairC'] if g['hair'] else None),(19,g['skin']),(20,g['eye']),(23,g['trunk']),(24,g['trim']),(25,g['topC'] if g['top'] else None),(27,g['gloveC']),(29,g['sole'])]:
    if hx:x=ws.cell(row=r,column=c);x.fill=PatternFill('solid',fgColor=hx.lstrip('#'));x.font=Font(name=F,size=10,color=textc(hx))
ws['B1'].comment=Comment('1 = the favourite of that tournament. You start against the lowest seed; each round brings a higher one.','South Paw')
ws['I1'].comment=Comment('From make_roster.py: 0 is weaker than the Easy CPU, 1 is past Hard. Sets reaction time, defence, accuracy and work rate.','South Paw')
ws['K1'].comment=Comment('Damage he takes is divided by this. In the Ultimate it also grows 7% each round (src/ui/career.js).','South Paw')
ws['L1'].comment=Comment('Damage he deals is multiplied by this. In the Ultimate it also grows 3.5% each round.','South Paw')
ws['M1'].comment=Comment('Share of normal damage he takes at the start of the tournament: 1 / Toughness.','South Paw')
ws.auto_filter.ref=f'A1:{get_column_letter(len(cols))}{len(rows)+1}'
for r in range(2,len(rows)+2):ws.row_dimensions[r].height=28
# ---- styles
s2=wb.create_sheet('Styles');head(s2,['Style','How he fights','Likes to fight at','Favourite game plan','Throws most','Changes to his skill','Fighters'],[17,52,16,18,30,40,9]);s2.freeze_panes='B2'
PLAN={'box':'Boxing at range','pressure':'Pressure','counter':'Countering','body':'Body work'};MODN={'aggr':'work rate','defend':'defence','slip':'dodging','combo':'combinations','tmin':'pause between moves','tmax':'pause between moves','react':'reaction time','punish':'punishing misses','acc':'reading punches','taunt':'taunting'}
for r,(k,st) in enumerate(ST.items(),2):
  rng=st.get('range',0);dist='Long range' if rng>0.12 else 'Medium-long' if rng>0.04 else 'In close' if rng<-0.1 else 'Medium-close' if rng<-0.03 else 'Middle distance'
  plan=', '.join(PLAN[p] for p,v in st.get('plan',{}).items() if v>=1) or 'Adapts to you'
  mix=', '.join(f'{p} x{v:g}' for p,v in sorted(st.get('mix',{}).items(),key=lambda kv:-kv[1]) if v>1) or 'Balanced'
  mods=[];seen=set()
  for m,v in st.get('mods',{}).items():
    nm=MODN[m]
    if nm in seen:continue
    seen.add(nm);mods.append(f"{nm} {'+' if v>1 else '-'}{abs(round((v-1)*100))}%")
  if st.get('power'):mods.append(f"power x{st['power']}")
  if st.get('tough'):mods.append(f"toughness x{st['tough']}")
  for c,v in enumerate([st['name'],st['desc'],dist,plan,mix,', '.join(mods),f'=COUNTIF(Roster!F:F,A{r})'],1):
    x=s2.cell(row=r,column=c,value=v);x.font=BF;x.border=BOX;x.alignment=Alignment(vertical='top',wrap_text=True)
  s2.row_dimensions[r].height=42
# ---- tournaments
s3=wb.create_sheet('Tournaments');head(s3,['Tournament','Short name','Fighters','Rounds','Opponents','Weakest skill','Strongest skill','Toughness from','Toughness to','Toughness growth per round'],[32,12,9,8,11,13,14,14,12,16]);s3.freeze_panes='B2'
for r,t in enumerate(d['tours'],2):
  vals=[t['name'],t['short'],t['size'],f'=LOG(C{r},2)',f'=COUNTIF(Roster!A:A,B{r})',f'=_xlfn.MINIFS(Roster!I:I,Roster!A:A,B{r})',f'=_xlfn.MAXIFS(Roster!I:I,Roster!A:A,B{r})',
        f'=_xlfn.MINIFS(Roster!K:K,Roster!A:A,B{r})',f'=_xlfn.MAXIFS(Roster!K:K,Roster!A:A,B{r})',0.07 if t['id']=='ultimate' else 0]
  for c,v in enumerate(vals,1):
    x=s3.cell(row=r,column=c,value=v);x.font=IN if c in (3,10) else BF;x.border=BOX
  for c in (6,7):s3.cell(row=r,column=c).number_format='0.00'
  for c in (8,9):s3.cell(row=r,column=c).number_format='0.00"x"'
  s3.cell(row=r,column=10).number_format='0%'
s3['J1'].comment=Comment('Set in src/ui/career.js (carGrit): Ultimate fighters take 7% less damage and hit 3.5% harder for each round you have won.','South Paw')
s3.cell(row=6,column=1,value='You make up the last place in every draw: Fighters = Opponents + 1. Blue numbers are inputs; the rest are formulas.').font=Font(name=F,size=9,italic=True,color='666666')
wb.save(os.path.join(HERE,'roster.xlsx'));print('saved')
